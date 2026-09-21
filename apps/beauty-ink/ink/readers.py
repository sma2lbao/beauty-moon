"""文件读取层：把 .md / .txt / .pdf / .docx 统一转为 LlamaIndex Document。

metadata 约定（所有格式共有）：
- source：文件绝对路径（增量更新 / 引用溯源 / 目录过滤的锚点）
- file_name：文件名（引用展示用）
- file_hash：文件内容 SHA-256（增量更新判断"是否变化"的依据）
- mtime：文件修改时间（ISO 8601，UTC）

PDF 额外逐页记录：page（1 起始页码）、total_pages（总页数），
供 M3 引用溯源时展示"文件名 + 页码"。
"""

import hashlib
from datetime import UTC, datetime
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from llama_index.core import Document

from .config import Settings

# 受支持的文档扩展名（M2：四种格式）
SUPPORTED_EXTS: frozenset[str] = frozenset({".md", ".txt", ".pdf", ".docx"})

# 通用 metadata 键：切分后的节点会继承
_COMMON_EXCLUDED_EMBED = ["source", "file_name", "file_hash", "mtime"]
_COMMON_EXCLUDED_LLM = ["file_hash", "mtime"]


def compute_file_hash(path: Path) -> str:
    """计算文件内容的 SHA-256，作为"文件是否变化"的判断依据。"""
    sha = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            sha.update(block)
    return sha.hexdigest()


def iter_supported_files(root: Path) -> list[Path]:
    """递归收集目录下所有受支持扩展名的文件。

    跳过隐藏目录（如 .venv、.git）与隐藏文件。
    """
    root = root.expanduser().resolve()
    if root.is_file():
        return [root] if root.suffix.lower() in SUPPORTED_EXTS else []
    files: list[Path] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        # 跳过任何路径段以点开头的文件（隐藏文件 / 隐藏目录内文件）
        if any(part.startswith(".") for part in path.relative_to(root).parts):
            continue
        if path.suffix.lower() in SUPPORTED_EXTS:
            files.append(path)
    return files


def _base_metadata(path: Path) -> dict[str, str]:
    """构造所有格式共有的 metadata（source / file_name / file_hash / mtime）。"""
    stat = path.stat()
    return {
        "source": str(path),
        "file_name": path.name,
        "file_hash": compute_file_hash(path),
        "mtime": datetime.fromtimestamp(stat.st_mtime, tz=UTC).isoformat(),
    }


def _read_markdown(path: Path) -> list[Document]:
    """读取 Markdown / 纯文本：单个 Document，整体交由上层切分器处理。"""
    return [
        Document(
            text=path.read_text(encoding="utf-8"),
            metadata=_base_metadata(path),
            excluded_embed_metadata_keys=_COMMON_EXCLUDED_EMBED,
            excluded_llm_metadata_keys=_COMMON_EXCLUDED_LLM,
        )
    ]


def _read_pdf(path: Path) -> list[Document]:
    """读取 PDF（PyMuPDF）：每一页生成一个 Document，并记录页码 metadata。

    逐页拆分的原因：切分后的文本块能继承 page 页码，
    引用溯源时可以精确到"文件名 + 第 N 页"。
    """
    docs: list[Document] = []
    with pymupdf.open(path) as pdf:
        total = pdf.page_count
        for i, page in enumerate(pdf, start=1):
            meta = _base_metadata(path) | {"page": str(i), "total_pages": str(total)}
            docs.append(
                Document(
                    text=page.get_text(),
                    metadata=meta,
                    excluded_embed_metadata_keys=_COMMON_EXCLUDED_EMBED + ["page", "total_pages"],
                    excluded_llm_metadata_keys=_COMMON_EXCLUDED_LLM + ["total_pages"],
                )
            )
    return docs


def _read_docx(path: Path) -> list[Document]:
    """读取 docx（python-docx）：按段落拼接为单个 Document（空段落跳过）。"""
    docx = DocxDocument(str(path))
    paragraphs = [p.text for p in docx.paragraphs if p.text.strip()]
    return [
        Document(
            text="\n".join(paragraphs),
            metadata=_base_metadata(path),
            excluded_embed_metadata_keys=_COMMON_EXCLUDED_EMBED,
            excluded_llm_metadata_keys=_COMMON_EXCLUDED_LLM,
        )
    ]


def read_documents(path: Path) -> list[Document]:
    """按扩展名分发读取：文件 → Document 列表（PDF 为多页多 Document）。"""
    ext = path.suffix.lower()
    if ext in {".md", ".txt"}:
        return _read_markdown(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    raise ValueError(f"不支持的文件类型：{path}（支持：{', '.join(sorted(SUPPORTED_EXTS))}）")


def load_embed_model(settings: Settings):
    """按 .env 中配置的模型名加载本地 HuggingFace Embedding（惰性导入避免拖慢 CLI）。"""
    from llama_index.embeddings.huggingface import HuggingFaceEmbedding

    if not settings.embedding_model:
        raise ValueError(
            "未配置 EMBEDDING_MODEL：请在 .env 中设置本地 Embedding 模型，"
            "例如 EMBEDDING_MODEL=BAAI/bge-m3（约 2GB）或 BAAI/bge-small-zh-v1.5（约 100MB）。"
            "注意：ingest 与 query 必须使用同一模型；更换模型后需删除 data/chroma 重新入库。"
        )
    return HuggingFaceEmbedding(model_name=settings.embedding_model)
